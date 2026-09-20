from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Pt


SOURCE = Path(r"C:\Users\slump\Desktop\软件开发框架\第1次课实践任务\实践任务单_第1次课.docx")
OUTPUT = Path(r"C:\Users\slump\Desktop\软件开发框架\第1次课实践任务\实践任务单_第1次课_已完成.docx")


def set_run_font(run, size=10.5, bold=False):
    run.font.name = "宋体"
    run.font.size = Pt(size)
    run.font.bold = bold
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:eastAsia"), "宋体")
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")


def write_cell(cell, text, size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    lines = text.split("\n")
    for index, line in enumerate(lines):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        paragraph.alignment = align
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.line_spacing = 1.1
        run = paragraph.add_run(line)
        set_run_font(run, size=size)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def add_paragraph_after(paragraph, text):
    new_element = OxmlElement("w:p")
    paragraph._p.addnext(new_element)
    new_paragraph = Paragraph(new_element, paragraph._parent)
    new_paragraph.paragraph_format.space_before = Pt(4)
    new_paragraph.paragraph_format.space_after = Pt(8)
    new_paragraph.paragraph_format.line_spacing = 1.25
    run = new_paragraph.add_run(text)
    set_run_font(run, size=10.5)
    return new_paragraph


def style_completed_cells(document):
    for table in document.tables:
        for row_index, row in enumerate(table.rows):
            if row_index == 0 and len(table.rows) > 1:
                for cell in row.cells:
                    shade_cell(cell, "D9EAF7")
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.font.bold = True
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(2)


def main():
    document = Document(SOURCE)
    tables = document.tables

    scenario_results = [
        ("借阅成功", "图书状态改为不可借，活动借阅数增加 1；默认短信通知和审计输出各执行一次。"),
        ("还书成功", "借阅记录标记为已归还，图书恢复可借状态，并输出一条还书审计记录。"),
        ("抛出 IllegalArgumentException", "提示 User not found；图书、借阅记录和通知均保持原状。"),
        ("抛出 IllegalStateException", "提示 Book unavailable；已有借阅记录保持不变，未重复发送通知。"),
    ]
    for row, (result, observation) in zip(tables[1].rows[1:], scenario_results):
        write_cell(row.cells[1], result)
        write_cell(row.cells[2], observation)

    runtime_map = (
        "Caller\n"
        "↓ CALL\n"
        "LibraryService.borrowBook()\n"
        "↓ READ\n"
        "InMemoryDatabase.findUser() / findBook()\n"
        "↓ DECIDE\n"
        "检查用户存在且图书可借\n"
        "↓ READ\n"
        "countActiveBorrowings()\n"
        "↓ DECIDE\n"
        "按用户类型比较借阅上限\n"
        "↓ WRITE\n"
        "BookCopy.borrow()\n"
        "↓ WRITE\n"
        "saveBorrowRecord()\n"
        "↓ EFFECT\n"
        "BorrowNotifier 通知并输出审计记录"
    )
    write_cell(tables[3].cell(0, 0), runtime_map, size=9.5)

    impact_values = [
        "2 个：LibraryService.java、Class01ChallengeTest.java。",
        "3 个：borrowBook() 及两个借阅上限测试。",
        "1 处业务分支。最终代码把两个上限集中在 borrowLimitFor() 的 switch 中。",
        "LibraryService 和对应测试需要识别 TEACHER；User 仍只保存通用 type 字段。",
        "新增并启用 2 个边界测试，原有 3 个回归测试继续保留。",
        "在 borrowLimitFor() 中加入 VIP 上限，并增加边界测试。借阅主流程无需另加分支。",
    ]
    for row, value in zip(tables[5].rows[1:], impact_values):
        write_cell(row.cells[1], value)

    smell_rows = [
        (
            "new LegacySmsClient()",
            "Coupling / Testability",
            "LibraryService 固定创建短信客户端，规则测试会连带执行短信副作用。",
            "接入 Email、消息队列或真实短信网关时，业务类随之修改。",
            "YES",
        ),
        (
            "new InMemoryDatabase()",
            "Dependency Direction",
            "业务类固定依赖内存存储，测试方便，但存储方式无法由外部装配。",
            "换成文件或关系数据库时，需要调整 LibraryService 的构造过程。",
            "UNSURE",
        ),
        (
            "borrowBook() 内只判断 STUDENT",
            "Change Locality",
            "V0 把上限规则写在借阅主流程中，新增 TEACHER 直接改动该方法。",
            "继续增加 VIP、访客等类型时，条件会继续增长。",
            "YES",
        ),
        (
            "borrowBook() 直接组织短信内容并发送",
            "Responsibility",
            "业务方法同时负责借阅规则和通知渠道细节，两个变化原因混在一起。",
            "增加 Email 或调整短信参数时，借阅主流程也要重新测试。",
            "YES",
        ),
        (
            "System.out.println(\"AUDIT...\")",
            "Responsibility",
            "审计输出位于业务方法中，目前只有一行且行为清楚。",
            "接入审计文件、操作人和时间信息后，代码量会明显增加。",
            "NO",
        ),
        (
            "borrowBook() 连续承担读取、判断、写入和通知",
            "Runtime Clarity",
            "执行顺序集中可见，阅读方便；继续加入规则后会变得拥挤。",
            "加入预约、欠费、库存锁定等规则时，主流程会拉长。",
            "UNSURE",
        ),
        (
            "User.type 使用字符串",
            "Cohesion / Type Safety",
            "拼写错误会落入默认处理，但当前只有两个已知类型且测试覆盖边界。",
            "用户类型增加并分散到多个模块后，错误更难定位。",
            "NO",
        ),
        (
            "returnBook() 依赖记录和图书同时存在",
            "Data Consistency",
            "当前内存库由同一服务写入，记录存在时通常也能找到图书。",
            "数据迁移或手工修复产生孤立记录时，可能出现空对象访问。",
            "NO",
        ),
    ]
    for row, values in zip(tables[6].rows[1:], smell_rows):
        for cell, value in zip(row.cells, values):
            write_cell(cell, value, size=8.5)

    refactor_values = [
        "LibraryService 内部直接创建 LegacySmsClient，规则测试也会执行短信发送。",
        "构造方法只接收 BorrowNotifier。App 装配 SmsBorrowNotifier；邮件场景传入 EmailBorrowNotifier，测试传入 Fake。",
        "Class01ChallengeTest 4 项通过，原有 LibraryServiceTest 3 项通过，共 7 项通过。",
        "替换通知渠道时，借阅上限、图书状态和借阅记录逻辑保持不动。",
        "新增 1 个接口、2 个实现和少量构造装配；Email 需要用户邮箱字段。",
    ]
    for cell, value in zip(tables[8].rows[1].cells, refactor_values):
        write_cell(cell, value, size=8.5)

    ai_reviews = [
        (
            "为通知能力提取接口并通过构造方法传入",
            "LibraryService 直接 new LegacySmsClient()",
            "规则测试可使用 Fake，新增 Email 只增加实现类",
        "增加接口和 App 装配代码",
            "ACCEPT。它直接解决测试副作用，也为第二个需求提供扩展点。",
        ),
        (
            "为每种用户类型建立借阅策略类",
            "V0 只有 STUDENT 一条上限判断",
            "用户类型大量增加时可分别维护规则",
            "当前要增加多个类和装配逻辑",
            "REJECT。现阶段只有 STUDENT 和 TEACHER 两个固定上限，集中 switch 更容易读。",
        ),
        (
            "把 InMemoryDatabase 也改为接口并注入",
            "LibraryService 直接 new InMemoryDatabase()",
            "将来替换数据库时业务类改动更少",
            "当前没有第二种存储实现，接口缺少真实变化依据",
            "MODIFY。记录为后续事项，等出现持久化需求时再处理。",
        ),
        (
            "把 User.type 从字符串改为枚举",
            "STUDENT 和 TEACHER 以字符串比较",
            "编译期可发现拼写错误",
            "需要修改模型、样例数据和相关测试",
            "MODIFY。方向合理，本次先用集中规则和边界测试控制范围。",
        ),
        (
            "把借阅成功事件发布到 EventBus，由监听器发送通知",
            "通知发生在借阅记录保存之后",
            "多个异步订阅者出现时能降低直接依赖",
            "引入事件模型、注册和失败处理，运行过程更难追踪",
            "REJECT。当前只有单一同步通知点，构造注入已满足需求。",
        ),
    ]
    for row, values in zip(tables[10].rows[1:], ai_reviews):
        for cell, value in zip(row.cells, values):
            write_cell(cell, value, size=8.3)

    email_analysis = [
        "V0 在 LibraryService 中直接创建并调用 LegacySmsClient，通知文本也在 borrowBook() 内拼接。",
        "借阅业务依赖具体短信实现，规则测试和通知渠道产生联动。",
        "新增 Email、接入真实网关或增加通知失败处理时，业务类会反复修改。",
        "保留原代码；增加 setter；构造注入 BorrowNotifier；使用事件总线。",
        "采用构造注入。依赖在对象创建时明确，App 继续装配短信实现，业务类只保留通知接口。",
        "收益是通知可替换，借阅主流程保持稳定。代价是新增接口、实现类和构造装配。",
        "Fake 测试证明规则验证无需短信客户端；Email 测试证明更换渠道时借阅判断没有变化。",
    ]
    for row, value in zip(tables[12].rows[1:], email_analysis):
        write_cell(row.cells[1], value, size=9.0)

    summary = (
        "我认为最危险的问题是 LibraryService 直接持有并调用具体短信客户端。"
        "它让借阅规则测试连带执行通知副作用，新增 Email 时也会触碰业务主流程。"
        "用户类型使用字符串同样存在拼写风险，但目前规则集中、测试覆盖边界，影响范围更容易控制。"
    )
    for paragraph in document.paragraphs:
        if paragraph.text.startswith("个人总结（不超过 200 字）"):
            add_paragraph_after(paragraph, summary)
            break

    style_completed_cells(document)
    document.core_properties.title = "第1次课课堂实践任务单已完成"
    document.core_properties.subject = "软件框架与设计模式实践"
    document.core_properties.comments = ""
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
